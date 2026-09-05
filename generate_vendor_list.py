#!/usr/bin/env python3
"""Generate a 100-row possible-vendor Excel list."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList
from collections import Counter

OUTPUT_PATH = "possible_vendors.xlsx"

VENDORS = [
    ("Aether Cloud Systems", "IT & Cloud", "Infrastructure", "United States", "Seattle", "Maya Chen", "Account Director", "Net 30", "USD", "Preferred", 4.8, 2014, "Multi-cloud landing zone partner"),
    ("Nordic Dataforge", "IT & Cloud", "Data Platform", "Sweden", "Stockholm", "Erik Lindqvist", "Sales Lead", "Net 30", "EUR", "Active", 4.6, 2016, "Snowflake and dbt specialists"),
    ("Kiteway Analytics", "IT & Cloud", "BI & Analytics", "United Kingdom", "London", "Priya Shah", "Customer Success", "Net 45", "GBP", "Active", 4.4, 2015, "Looker and Tableau implementations"),
    ("Harborline Logistics", "Logistics", "Freight Forwarding", "Netherlands", "Rotterdam", "Joost van Dijk", "Ops Manager", "Net 30", "EUR", "Preferred", 4.7, 2008, "EU port and last-mile coverage"),
    ("Pacific Rim Freight", "Logistics", "Ocean Freight", "Singapore", "Singapore", "Wei Ling Tan", "Trade Director", "Net 45", "USD", "Active", 4.3, 2001, "Asia-Pacific ocean lanes"),
    ("Alpine Cold Chain", "Logistics", "Cold Storage", "Switzerland", "Zurich", "Lena Meier", "Key Account Mgr", "Net 30", "CHF", "Preferred", 4.9, 2010, "Temperature-controlled warehousing"),
    ("Cedar & Oak Office", "Office Supplies", "Stationery", "Canada", "Toronto", "Noah Patel", "Sales Rep", "Net 15", "CAD", "Active", 4.1, 1998, "Bulk office consumables"),
    ("Paperfold Co.", "Office Supplies", "Print & Paper", "Germany", "Hamburg", "Anja Vogel", "Procurement Lead", "Net 30", "EUR", "Prospect", 3.9, 1992, "Sustainable paper products"),
    ("Brightspan Facilities", "Facilities", "Cleaning", "Ireland", "Dublin", "Ciaran Doyle", "Site Lead", "Net 15", "EUR", "Active", 4.2, 2007, "Night-shift cleaning crews"),
    ("Ironlatch Security", "Security", "Physical Security", "United States", "Chicago", "Denise Harper", "Regional Manager", "Net 30", "USD", "Active", 4.5, 2003, "On-site guards and access control"),
    ("Ciphergate Labs", "Security", "Cybersecurity", "Israel", "Tel Aviv", "Yael Cohen", "Solutions Architect", "Net 30", "USD", "Preferred", 4.8, 2017, "SOC and endpoint protection"),
    ("Lumenfield Telecom", "Telecom", "Connectivity", "France", "Paris", "Hugo Martin", "Enterprise AE", "Net 30", "EUR", "Active", 4.0, 2005, "MPLS and SD-WAN circuits"),
    ("Orbit Mobile Fleet", "Telecom", "Mobile Devices", "South Korea", "Seoul", "Min-jun Park", "Channel Partner", "Net 45", "USD", "Prospect", 3.8, 2012, "Corporate handset programs"),
    ("Quarry Peak Hardware", "Hardware", "Servers", "Taiwan", "Taipei", "Hsin-Yi Wu", "OEM Manager", "Net 45", "USD", "Active", 4.4, 2006, "Rack servers and spare parts"),
    ("Nimbus Printworks", "Hardware", "Peripherals", "Japan", "Osaka", "Kenji Sato", "Sales Engineer", "Net 30", "JPY", "Active", 4.1, 1999, "Printers and scanning devices"),
    ("Greenlane Packaging", "Manufacturing", "Packaging", "Poland", "Krakow", "Marta Nowak", "Export Manager", "Net 30", "EUR", "Preferred", 4.6, 2009, "Retail-ready corrugated packs"),
    ("Forge & Filament", "Manufacturing", "Metalwork", "Italy", "Milan", "Luca Bianchi", "Plant Liaison", "Net 45", "EUR", "Active", 4.2, 1987, "Custom metal fixtures"),
    ("Harvestfield Foods", "Food & Beverage", "Dry Goods", "Spain", "Valencia", "Carmen Ruiz", "Category Lead", "Net 15", "EUR", "Preferred", 4.7, 1995, "Private-label pantry staples"),
    ("Northwind Dairy Co.", "Food & Beverage", "Dairy", "Denmark", "Aarhus", "Sofie Berg", "Export AE", "Net 15", "EUR", "Active", 4.5, 1978, "UHT and chilled dairy"),
    ("Cinder Roast Coffee", "Food & Beverage", "Beverages", "Colombia", "Medellin", "Andres Gil", "Trade Manager", "Net 30", "USD", "Active", 4.6, 2011, "Specialty coffee supply"),
    ("Blueharbor Seafood", "Food & Beverage", "Fresh Protein", "Norway", "Bergen", "Ingrid Dahl", "Sales Director", "Net 7", "EUR", "Preferred", 4.8, 2004, "Traceable salmon and whitefish"),
    ("Suncrest Produce", "Food & Beverage", "Fresh Produce", "Mexico", "Guadalajara", "Elena Vargas", "Grower Rep", "Net 7", "USD", "Active", 4.3, 2002, "Seasonal fruit programs"),
    ("Helix Staffing Group", "HR & Staffing", "Contractors", "United States", "Austin", "Jordan Miles", "Talent Partner", "Net 15", "USD", "Active", 4.2, 2013, "Data and engineering contractors"),
    ("Pinnacle People Ops", "HR & Staffing", "RPO", "India", "Bengaluru", "Ananya Iyer", "Delivery Lead", "Net 30", "USD", "Prospect", 4.0, 2016, "Volume hiring for IT roles"),
    ("Meridian Legal Partners", "Legal", "Commercial Law", "United Kingdom", "Manchester", "Owen Clarke", "Partner", "Upon receipt", "GBP", "Active", 4.6, 2000, "Vendor contracts and IP"),
    ("Oakridge Audit LLC", "Finance", "Audit", "United States", "Boston", "Rachel Kim", "Engagement Mgr", "Net 30", "USD", "Active", 4.5, 1994, "SOC 2 and vendor audits"),
    ("Ledgerbay Payments", "Finance", "Payments", "Ireland", "Cork", "Niamh Walsh", "Partnerships", "Net 15", "EUR", "Preferred", 4.7, 2018, "Cross-border payout rails"),
    ("Canvas & Click Media", "Marketing", "Digital Ads", "Australia", "Sydney", "Tom Nguyen", "Media Director", "Net 30", "AUD", "Active", 4.1, 2014, "Performance media buying"),
    ("Storyline Creative", "Marketing", "Brand", "Brazil", "Sao Paulo", "Beatriz Lima", "Creative Lead", "Net 30", "BRL", "Prospect", 3.9, 2019, "Campaign and packaging design"),
    ("Summit Advisory Group", "Consulting", "Strategy", "United States", "New York", "Alex Rivera", "Principal", "Net 15", "USD", "Active", 4.4, 2008, "Operating model design"),
    ("Bytebridge Consulting", "Consulting", "Data Engineering", "Germany", "Berlin", "Jonas Keller", "Practice Lead", "Net 30", "EUR", "Preferred", 4.8, 2015, "Kafka and lakehouse builds"),
    ("Clearpath ERP", "IT & Cloud", "ERP", "Netherlands", "Amsterdam", "Sanne de Boer", "Solution Consultant", "Net 45", "EUR", "Active", 4.3, 2006, "SAP S/4 rollout support"),
    ("Redwood Identity", "IT & Cloud", "IAM", "Canada", "Vancouver", "Chris Okonkwo", "Security AE", "Net 30", "CAD", "Active", 4.5, 2017, "SSO and privileged access"),
    ("Falcon Eye Monitoring", "IT & Cloud", "Observability", "United States", "Denver", "Sam Ortega", "TAM", "Net 30", "USD", "Preferred", 4.6, 2016, "APM and log pipelines"),
    ("Driftwood APIs", "IT & Cloud", "Integration", "Estonia", "Tallinn", "Liis Kask", "Product Partner", "Net 30", "EUR", "Prospect", 4.0, 2020, "iPaaS connectors"),
    ("Copperline Electrical", "Facilities", "Electrical", "United Kingdom", "Birmingham", "Harry Cole", "Contracts Mgr", "Net 30", "GBP", "Active", 4.2, 1991, "Fit-out and maintenance"),
    ("Stonewell HVAC", "Facilities", "Climate Control", "United States", "Atlanta", "Monica Reed", "Branch Manager", "Net 30", "USD", "Active", 4.1, 1985, "Warehouse HVAC service"),
    ("Evergreen Waste Co.", "Facilities", "Waste Management", "Sweden", "Gothenburg", "Freja Holm", "Account Mgr", "Net 30", "SEK", "Active", 4.3, 2003, "Recycling and composting"),
    ("Skyline Uniforms", "Apparel", "Workwear", "Portugal", "Porto", "Rui Costa", "Wholesale Lead", "Net 45", "EUR", "Prospect", 3.8, 2009, "Store associate uniforms"),
    ("Threadline Textiles", "Apparel", "Soft Goods", "Turkey", "Istanbul", "Elif Yilmaz", "Export AE", "Net 45", "USD", "Active", 4.2, 1997, "Private-label apparel"),
    ("Beacon POS Systems", "Hardware", "Point of Sale", "United States", "San Jose", "Kevin Tran", "Channel Mgr", "Net 30", "USD", "Preferred", 4.6, 2011, "Checkout terminals and PIN pads"),
    ("Quickscan Labels", "Hardware", "Barcode", "China", "Shenzhen", "Liu Fang", "OEM Sales", "Net 45", "USD", "Active", 4.0, 2005, "Scanners and label printers"),
    ("Atlas Fleet Services", "Logistics", "Fleet Maintenance", "Germany", "Munich", "Stefan Huber", "Fleet Lead", "Net 15", "EUR", "Active", 4.3, 2000, "Light commercial vehicle service"),
    ("Riverbend Pallet Co.", "Manufacturing", "Pallets", "United States", "Memphis", "Dale Cooper", "Plant Manager", "Net 15", "USD", "Active", 4.1, 1989, "CHEP-compatible pallets"),
    ("Glasshouse Bottling", "Manufacturing", "Containers", "Czechia", "Prague", "Petra Novak", "Key Accounts", "Net 30", "EUR", "Prospect", 3.9, 2004, "Glass and PET bottles"),
    ("Silverpine Chemicals", "Manufacturing", "Cleaning Agents", "Belgium", "Antwerp", "Pieter Jans", "Chem Sales", "Net 30", "EUR", "Active", 4.2, 1996, "Food-safe sanitizers"),
    ("Nova Lab Supplies", "Lab & QA", "Testing Kits", "Switzerland", "Basel", "Clara Hofmann", "Scientific AE", "Net 30", "CHF", "Active", 4.5, 2012, "Food safety test kits"),
    ("Precision Scale Co.", "Lab & QA", "Metrology", "Japan", "Tokyo", "Yuki Tanaka", "Export Engineer", "Net 45", "JPY", "Preferred", 4.7, 1982, "Certified weighing equipment"),
    ("Horizon Insurance Brokers", "Finance", "Insurance", "United Kingdom", "Edinburgh", "Fiona Grant", "Broker", "Upon receipt", "GBP", "Active", 4.4, 1990, "Cargo and liability cover"),
    ("Waypoint Travel Desk", "Travel", "Corporate Travel", "United States", "Dallas", "Luis Romero", "Travel Advisor", "Net 15", "USD", "Active", 4.0, 2007, "Duty of care and bookings"),
    ("Larkspur Catering", "Food & Beverage", "Catering", "France", "Lyon", "Camille Petit", "Events Lead", "Net 7", "EUR", "Prospect", 4.1, 2013, "Office and event catering"),
    ("Moss & Pine Landscaping", "Facilities", "Grounds", "New Zealand", "Auckland", "Hayley Scott", "Site Supervisor", "Net 15", "NZD", "Active", 4.2, 2006, "Storefront landscaping"),
    ("Arclight Signage", "Marketing", "In-store Media", "United States", "Los Angeles", "Devon Brooks", "Production AE", "Net 30", "USD", "Active", 4.3, 2008, "Digital and printed signage"),
    ("Pixelnest Studios", "Marketing", "Content", "Poland", "Warsaw", "Kasia Wozniak", "Producer", "Net 30", "EUR", "Prospect", 3.9, 2018, "Product photography"),
    ("Tidewatch Compliance", "Legal", "Regulatory", "Germany", "Frankfurt", "Helena Fuchs", "Counsel", "Net 15", "EUR", "Active", 4.6, 2010, "GDPR and food labeling"),
    ("Northstar Translation", "Professional Services", "Localization", "Spain", "Barcelona", "Marc Sole", "PM", "Net 15", "EUR", "Active", 4.3, 2014, "Pack copy in 20+ languages"),
    ("Kinetic Training Lab", "HR & Staffing", "L&D", "United States", "Minneapolis", "Aisha Brown", "Learning Partner", "Net 30", "USD", "Prospect", 4.0, 2017, "Store ops e-learning"),
    ("Cobalt Print House", "Office Supplies", "Commercial Print", "United Kingdom", "Leeds", "Ben Walsh", "Print Broker", "Net 15", "GBP", "Active", 4.1, 2001, "Catalogs and circulars"),
    ("Willowbrook Furniture", "Facilities", "Fixtures", "Vietnam", "Ho Chi Minh City", "Tran Minh", "Export Manager", "Net 45", "USD", "Active", 4.2, 2005, "Store fixtures and shelving"),
    ("Summit Rack Systems", "Facilities", "Warehouse Racking", "South Africa", "Johannesburg", "Thabo Nkosi", "Projects Lead", "Net 45", "ZAR", "Preferred", 4.5, 1999, "Pallet racking installs"),
    ("Ion Battery Works", "Hardware", "Power", "South Korea", "Busan", "Ji-woo Han", "OEM AE", "Net 45", "USD", "Prospect", 4.0, 2015, "UPS and backup batteries"),
    ("Coralnet ISP", "Telecom", "Internet", "Philippines", "Manila", "Isabel Cruz", "Enterprise AE", "Net 30", "USD", "Active", 3.8, 2009, "Store broadband links"),
    ("Mapleleaf Janitorial", "Facilities", "Janitorial", "Canada", "Calgary", "Ryan Fraser", "Ops Coordinator", "Net 15", "CAD", "Active", 4.0, 2002, "Regional store cleaning"),
    ("Vanguard Pest Control", "Facilities", "Pest Control", "Australia", "Melbourne", "Olivia Grant", "Technician Lead", "Net 15", "AUD", "Active", 4.3, 1993, "Food-retail pest programs"),
    ("Solara Energy Co.", "Utilities", "Renewables", "Spain", "Madrid", "Diego Alvarez", "Energy Analyst", "Net 30", "EUR", "Preferred", 4.6, 2013, "On-site solar PPAs"),
    ("Gridline Utilities", "Utilities", "Electricity", "United States", "Houston", "Pat Simmons", "Account Exec", "Net 15", "USD", "Active", 4.1, 1980, "Deregulated power supply"),
    ("Aquaflow Filtration", "Utilities", "Water Treatment", "Israel", "Haifa", "Noam Levi", "Engineer AE", "Net 30", "USD", "Prospect", 4.2, 2011, "Store water filtration"),
    ("Brightpath Lighting", "Hardware", "Lighting", "China", "Hangzhou", "Chen Wei", "Export Sales", "Net 45", "USD", "Active", 4.0, 2008, "LED retail lighting"),
    ("Ember Fire Safety", "Security", "Fire Protection", "United Kingdom", "Glasgow", "Iain Murray", "Compliance Lead", "Net 30", "GBP", "Active", 4.5, 1988, "Sprinklers and extinguishers"),
    ("Lockridge Keys", "Security", "Access Hardware", "United States", "Phoenix", "Tara Benson", "Branch Lead", "Net 15", "USD", "Active", 4.1, 1996, "Locks, keys, and safes"),
    ("Helios Solar Glass", "Manufacturing", "Building Materials", "India", "Pune", "Rohit Desai", "Project AE", "Net 45", "INR", "Prospect", 3.9, 2014, "Facade and canopy glass"),
    ("Cinderblock Construction", "Facilities", "Fit-out", "United Arab Emirates", "Dubai", "Omar Haddad", "Project Mgr", "Milestone", "AED", "Active", 4.2, 2007, "Store build-outs"),
    ("Northwind Snow & Ice", "Facilities", "Seasonal Services", "Canada", "Montreal", "Julie Gagnon", "Dispatch Lead", "Net 15", "CAD", "Seasonal", 4.0, 2004, "Parking-lot snow removal"),
    ("Pebble Path Parking", "Facilities", "Parking Ops", "United States", "San Francisco", "Nate Feldman", "Ops AE", "Net 30", "USD", "Prospect", 3.7, 2016, "Lot management software"),
    ("Asterisk Helpdesk", "IT & Cloud", "ITSM", "India", "Hyderabad", "Neha Reddy", "Service Lead", "Net 30", "USD", "Active", 4.3, 2012, "L1/L2 service desk"),
    ("Cobalt Backup Vault", "IT & Cloud", "Backup", "Finland", "Helsinki", "Mikko Virtanen", "Solutions Eng", "Net 30", "EUR", "Preferred", 4.7, 2015, "Immutable backups"),
    ("Sable CRM Works", "IT & Cloud", "CRM", "United States", "Philadelphia", "Gina Rossi", "Implementation", "Net 30", "USD", "Active", 4.2, 2010, "Salesforce customizations"),
    ("Windmill GIS", "IT & Cloud", "Geospatial", "Netherlands", "Utrecht", "Bram Visser", "GIS Consultant", "Net 30", "EUR", "Prospect", 4.1, 2018, "Store location analytics"),
    ("Harbor Tax Advisors", "Finance", "Tax", "Luxembourg", "Luxembourg", "Elise Weber", "Tax Director", "Upon receipt", "EUR", "Active", 4.5, 2001, "Cross-border VAT"),
    ("Quill & Bind Records", "Professional Services", "Records Mgmt", "United States", "Kansas City", "Holly Price", "Account Mgr", "Net 30", "USD", "Active", 4.0, 1992, "Offsite document storage"),
    ("Nightowl Courier", "Logistics", "Last Mile", "United Kingdom", "London", "Jamal Adeyemi", "Ops Director", "Net 7", "GBP", "Preferred", 4.6, 2016, "Same-day urban delivery"),
    ("Overland Rail Cargo", "Logistics", "Rail", "United States", "Chicago", "Bill Keene", "Intermodal AE", "Net 30", "USD", "Active", 4.2, 1984, "Domestic rail containers"),
    ("Skylark Air Cargo", "Logistics", "Air Freight", "United Arab Emirates", "Abu Dhabi", "Fatima Al Suwaidi", "Cargo AE", "Net 15", "AED", "Active", 4.3, 2006, "Priority perishable airlift"),
    ("Tundra Insulation", "Manufacturing", "Insulation", "Finland", "Oulu", "Aino Laakso", "Sales Eng", "Net 45", "EUR", "Prospect", 4.0, 2003, "Cold-room insulation panels"),
    ("Citrus Grove Oils", "Food & Beverage", "Ingredients", "Italy", "Bari", "Giulia Greco", "Export Lead", "Net 30", "EUR", "Active", 4.4, 1998, "Olive oil and extracts"),
    ("Ambergrain Mills", "Food & Beverage", "Bakery", "France", "Lille", "Paul Renard", "Wholesale AE", "Net 15", "EUR", "Active", 4.3, 1975, "Flour and bakery mixes"),
    ("Saltmarsh Spices", "Food & Beverage", "Spices", "India", "Kochi", "Meera Nair", "Commodity AE", "Net 30", "USD", "Preferred", 4.7, 1986, "Traceable spice lots"),
    ("Polar Ice Works", "Food & Beverage", "Ice & Refrigerants", "Canada", "Winnipeg", "Kyle McLeod", "Route Mgr", "Net 7", "CAD", "Active", 4.1, 2000, "Packaged ice supply"),
    ("Lumen HR Benefits", "HR & Staffing", "Benefits", "United States", "Hartford", "Ellen Cho", "Benefits Consultant", "Net 30", "USD", "Active", 4.2, 2009, "Health and wellness plans"),
    ("Fieldnote Survey Co.", "Professional Services", "Research", "Netherlands", "The Hague", "Eva Bakker", "Research Lead", "Net 30", "EUR", "Prospect", 4.0, 2015, "Shopper insight studies"),
    ("Copperleaf Recycling", "Facilities", "Recycling", "Germany", "Cologne", "Tim Bauer", "Materials AE", "Net 30", "EUR", "Active", 4.4, 2008, "Cardboard and plastics take-back"),
    ("Ironclad Uniform Laundry", "Facilities", "Laundry", "United States", "Cincinnati", "Pam Ellis", "Route Supervisor", "Net 15", "USD", "Active", 4.1, 1994, "Workwear rental and wash"),
    ("Sagebrush Landscaping", "Facilities", "Irrigation", "United States", "Phoenix", "Cody Hale", "Field Lead", "Net 15", "USD", "Seasonal", 3.8, 2011, "Xeriscape maintenance"),
    ("Nimbus Drone Inspect", "Professional Services", "Inspection", "Australia", "Brisbane", "Mia Taylor", "Pilot Lead", "Net 15", "AUD", "Prospect", 4.2, 2019, "Roof and yard drone surveys"),
    ("Keystone Pallet Wrap", "Manufacturing", "Film & Wrap", "Malaysia", "Penang", "Aisha Rahman", "Export AE", "Net 45", "USD", "Active", 4.0, 2007, "Stretch film and banding"),
    ("Horizon Badge Co.", "Office Supplies", "ID & Badges", "United States", "Portland", "Lee Nakamura", "Sales Coord", "Net 15", "USD", "Active", 4.1, 2010, "Employee badges and lanyards"),
    ("Bluecedar Ergonomics", "Office Supplies", "Furniture", "Sweden", "Malmo", "Oskar Berg", "Workplace AE", "Net 30", "SEK", "Prospect", 4.3, 2016, "Ergonomic office chairs"),
    ("Atlas Language Line", "Professional Services", "Interpreting", "United States", "Miami", "Sofia Alvarez", "Vendor Mgr", "Net 15", "USD", "Active", 4.4, 2005, "On-demand interpreting"),
    ("Pioneer Safety Gear", "Apparel", "PPE", "Germany", "Dortmund", "Lukas Brandt", "Safety AE", "Net 30", "EUR", "Preferred", 4.6, 1991, "Gloves, helmets, high-vis"),
    ("Starlight Event Hire", "Marketing", "Events", "United Kingdom", "Bristol", "Chloe Adams", "Hire Manager", "Net 7", "GBP", "Seasonal", 4.0, 2012, "Pop-up store event kit"),
]


def slug_email(name: str, company: str) -> str:
    first, last = name.lower().split()[:2]
    domain = (
        company.lower()
        .replace("&", "and")
        .replace(".", "")
        .replace(",", "")
        .replace("'", "")
    )
    domain = "".join(ch if ch.isalnum() else "-" for ch in domain)
    domain = "-".join(part for part in domain.split("-") if part)
    return f"{first}.{last}@{domain}.example"


def phone_for(country: str, index: int) -> str:
    prefixes = {
        "United States": "+1-312",
        "Canada": "+1-416",
        "United Kingdom": "+44-20",
        "Germany": "+49-30",
        "France": "+33-1",
        "Netherlands": "+31-20",
        "Sweden": "+46-8",
        "Switzerland": "+41-44",
        "Ireland": "+353-1",
        "Israel": "+972-3",
        "South Korea": "+82-2",
        "Taiwan": "+886-2",
        "Japan": "+81-3",
        "Poland": "+48-12",
        "Italy": "+39-02",
        "Spain": "+34-91",
        "Denmark": "+45-32",
        "Colombia": "+57-4",
        "Norway": "+47-55",
        "Mexico": "+52-33",
        "India": "+91-80",
        "Australia": "+61-2",
        "Brazil": "+55-11",
        "Estonia": "+372-6",
        "Portugal": "+351-22",
        "Turkey": "+90-212",
        "China": "+86-755",
        "Czechia": "+420-2",
        "Belgium": "+32-3",
        "New Zealand": "+64-9",
        "Vietnam": "+84-28",
        "South Africa": "+27-11",
        "Philippines": "+63-2",
        "United Arab Emirates": "+971-4",
        "Finland": "+358-9",
        "Luxembourg": "+352-2",
        "Malaysia": "+60-4",
    }
    prefix = prefixes.get(country, "+1-202")
    return f"{prefix}-{4000 + index:04d}-{100 + index:04d}"


def website_for(company: str) -> str:
    slug = "".join(ch.lower() if ch.isalnum() else "" for ch in company.split()[0])
    return f"https://www.{slug}-vendor.example"


def main() -> None:
    assert len(VENDORS) == 100, f"expected 100 vendors, got {len(VENDORS)}"

    wb = Workbook()
    ws = wb.active
    ws.title = "Vendors"

    headers = [
        "Vendor ID",
        "Vendor Name",
        "Category",
        "Subcategory",
        "Country",
        "City",
        "Contact Person",
        "Job Title",
        "Email",
        "Phone",
        "Website",
        "Payment Terms",
        "Currency",
        "Status",
        "Rating",
        "Year Established",
        "Notes",
    ]
    ws.append(headers)

    for i, row in enumerate(VENDORS, start=1):
        (
            name,
            category,
            subcategory,
            country,
            city,
            contact,
            title,
            terms,
            currency,
            status,
            rating,
            year,
            notes,
        ) = row
        ws.append(
            [
                f"VND-{i:03d}",
                name,
                category,
                subcategory,
                country,
                city,
                contact,
                title,
                slug_email(contact, name),
                phone_for(country, i),
                website_for(name),
                terms,
                currency,
                status,
                rating,
                year,
                notes,
            ]
        )

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    alt_fill = PatternFill("solid", fgColor="F2F7FB")
    wrap = Alignment(vertical="center", wrap_text=True)
    center = Alignment(vertical="center", horizontal="center")

    status_fills = {
        "Preferred": PatternFill("solid", fgColor="C6EFCE"),
        "Active": PatternFill("solid", fgColor="DDEBF7"),
        "Prospect": PatternFill("solid", fgColor="FFF2CC"),
        "Seasonal": PatternFill("solid", fgColor="FCE4D6"),
    }

    for col in range(1, len(headers) + 1):
        cell = ws.cell(1, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=2, max_row=101, max_col=len(headers)):
        for cell in row:
            cell.font = Font(name="Calibri", size=11)
            cell.border = thin
            cell.alignment = wrap
            if cell.row % 2 == 0:
                cell.fill = alt_fill
        status_cell = row[13]
        status_cell.fill = status_fills.get(status_cell.value, status_cell.fill)
        status_cell.alignment = center
        row[0].alignment = center
        row[14].alignment = center
        row[14].number_format = "0.0"
        row[15].alignment = center

    ws.conditional_formatting.add(
        "O2:O101",
        ColorScaleRule(
            start_type="num",
            start_value=3.5,
            start_color="F8696B",
            mid_type="num",
            mid_value=4.2,
            mid_color="FFEB84",
            end_type="num",
            end_value=5.0,
            end_color="63BE7B",
        ),
    )

    table = Table(displayName="VendorList", ref="A1:Q101")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)

    widths = {
        "A": 12,
        "B": 28,
        "C": 22,
        "D": 20,
        "E": 22,
        "F": 18,
        "G": 18,
        "H": 20,
        "I": 42,
        "J": 20,
        "K": 34,
        "L": 16,
        "M": 12,
        "N": 12,
        "O": 10,
        "P": 18,
        "Q": 36,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.auto_filter.ref = "A1:Q101"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"

    summary = wb.create_sheet("Summary")
    summary["A1"] = "Possible vendors — summary"
    summary["A1"].font = Font(name="Calibri", bold=True, size=16, color="1F4E79")
    summary.merge_cells("A1:C1")

    summary["A3"] = "Total vendors"
    summary["B3"] = 100
    summary["A4"] = "Preferred vendors"
    summary["B4"] = sum(1 for v in VENDORS if v[9] == "Preferred")
    summary["A5"] = "Active vendors"
    summary["B5"] = sum(1 for v in VENDORS if v[9] == "Active")
    summary["A6"] = "Prospect vendors"
    summary["B6"] = sum(1 for v in VENDORS if v[9] == "Prospect")
    summary["A7"] = "Seasonal vendors"
    summary["B7"] = sum(1 for v in VENDORS if v[9] == "Seasonal")
    summary["A8"] = "Average rating"
    summary["B8"] = round(sum(v[10] for v in VENDORS) / len(VENDORS), 2)
    summary["B8"].number_format = "0.00"

    for r in range(3, 9):
        summary[f"A{r}"].font = Font(name="Calibri", bold=True, size=11)
        summary[f"B{r}"].font = Font(name="Calibri", size=11)

    summary["A10"] = "Vendors by category"
    summary["A10"].font = Font(name="Calibri", bold=True, size=13, color="1F4E79")
    summary["A11"] = "Category"
    summary["B11"] = "Count"
    for cell in (summary["A11"], summary["B11"]):
        cell.fill = header_fill
        cell.font = header_font

    counts = Counter(v[1] for v in VENDORS)
    for i, (category, count) in enumerate(sorted(counts.items(), key=lambda x: (-x[1], x[0])), start=12):
        summary[f"A{i}"] = category
        summary[f"B{i}"] = count

    last_cat_row = 11 + len(counts)
    pie = PieChart()
    pie.title = "Vendors by category"
    labels = Reference(summary, min_col=1, min_row=12, max_row=last_cat_row)
    data = Reference(summary, min_col=2, min_row=11, max_row=last_cat_row)
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(labels)
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.dataLabels.showVal = True
    pie.dataLabels.showCatName = False
    pie.width = 18
    pie.height = 10
    summary.add_chart(pie, "D10")

    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 12
    summary.column_dimensions["C"].width = 16

    notes = wb.create_sheet("Readme")
    notes["A1"] = "Possible vendor list"
    notes["A1"].font = Font(name="Calibri", bold=True, size=16, color="1F4E79")
    notes.merge_cells("A1:B1")
    readme_rows = [
        ("Rows", "100 unique possible vendors (VND-001 to VND-100)"),
        ("Sheets", "Vendors (full list), Summary (counts and chart), Readme"),
        ("How to use", "Filter the Vendors table by Category, Country, or Status"),
        ("Emails / phones / sites", "Sample contact details using .example domains, not live accounts"),
        ("Statuses", "Preferred, Active, Prospect, Seasonal"),
        ("Ratings", "1.0–5.0 scale; higher is better"),
    ]
    notes["A3"] = "Field"
    notes["B3"] = "Description"
    notes["A3"].fill = header_fill
    notes["B3"].fill = header_fill
    notes["A3"].font = header_font
    notes["B3"].font = header_font
    for i, (field, desc) in enumerate(readme_rows, start=4):
        notes[f"A{i}"] = field
        notes[f"B{i}"] = desc
        notes[f"A{i}"].font = Font(name="Calibri", bold=True, size=11)
        notes[f"B{i}"].font = Font(name="Calibri", size=11)
        notes[f"A{i}"].alignment = wrap
        notes[f"B{i}"].alignment = wrap
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 78

    wb.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH} with {len(VENDORS)} vendors")


if __name__ == "__main__":
    main()
